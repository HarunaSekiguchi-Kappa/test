import copy
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def replace_placeholders_in_text(text: str, data: dict) -> str:
    """
    {{key}} を data[key] に置換する。
    None の場合は空文字にする。
    """
    if text is None:
        return text

    for key, value in data.items():
        placeholder = "{{" + key + "}}"
        replacement = "" if value is None else str(value)
        text = text.replace(placeholder, replacement)

    return text


def replace_in_text_frame(text_frame, data: dict) -> None:
    """
    テキストフレーム内のプレースホルダーを置換する。
    書式をなるべく残すため、run単位で置換する。
    """
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.text = replace_placeholders_in_text(run.text, data)


def process_shape(shape, data: dict) -> None:
    """
    スライド内の図形、テキストボックス、表、グループ図形を処理する。
    """
    if shape.has_text_frame:
        replace_in_text_frame(shape.text_frame, data)

    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                replace_in_text_frame(cell.text_frame, data)

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub_shape in shape.shapes:
            process_shape(sub_shape, data)


def duplicate_slide(prs: Presentation, slide):
    """
    スライドを同じプレゼン内に複製する。
    python-pptxには公式のduplicate機能がないため、XMLをコピーする。
    """
    blank_layout = prs.slide_layouts[6]
    new_slide = prs.slides.add_slide(blank_layout)

    # 空白スライドに最初から入っているshapeを削除
    for shape in list(new_slide.shapes):
        element = shape.element
        element.getparent().remove(element)

    # 元スライドのshapeをコピー
    for shape in slide.shapes:
        new_element = copy.deepcopy(shape.element)
        new_slide.shapes._spTree.insert_element_before(new_element, "p:extLst")

    # 背景をコピー
    if slide.background:
        new_slide.background._element = copy.deepcopy(slide.background._element)

    return new_slide


def remove_original_template_slides(prs: Presentation, count: int) -> None:
    """
    先頭 count 枚のテンプレートスライドを削除する。
    """
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    for i in range(count):
        xml_slides.remove(slides[i])


def fill_slide(slide, data: dict) -> None:
    for shape in slide.shapes:
        process_shape(shape, data)


def make_deck(template_path: str, json_path: str, output_path: str) -> None:
    template_path = Path(template_path)
    json_path = Path(json_path)
    output_path = Path(output_path)

    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    if not json_path.exists():
        raise FileNotFoundError(f"JSON not found: {json_path}")

    with open(json_path, "r", encoding="utf-8") as f:
        papers = json.load(f)

    if isinstance(papers, dict):
        papers = [papers]

    if not isinstance(papers, list):
        raise ValueError("JSON must be an object or a list of objects.")

    prs = Presentation(str(template_path))

    template_slide_count = len(prs.slides)
    if template_slide_count == 0:
        raise ValueError("Template has no slides.")

    template_slides = list(prs.slides)

    for index, paper in enumerate(papers, start=1):
        data = dict(paper)
        data.setdefault("paper_no", f"論文{index}")

        for template_slide in template_slides:
            new_slide = duplicate_slide(prs, template_slide)
            fill_slide(new_slide, data)

    remove_original_template_slides(prs, template_slide_count)

    prs.save(str(output_path))
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("Usage:")
        print("  python make_paper_slides.py template.pptx papers.json output.pptx")
        sys.exit(1)

    make_deck(
        template_path=sys.argv[1],
        json_path=sys.argv[2],
        output_path=sys.argv[3],
    )